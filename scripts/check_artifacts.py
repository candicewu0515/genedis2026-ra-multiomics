"""Check that the final figures, report and slides match after the documents are built."""
from pathlib import Path
import hashlib,json,sys,zipfile
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
r=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else Path(__file__).resolve().parents[1];fig=r/'results/figures/main';p=r/'presentation';checks=[]
names=['Fig1','Fig2','Fig3','Fig4','ExtData_Fig1','ExtData_Fig2']
expected={hashlib.sha256((fig/(n+'.png')).read_bytes()).hexdigest() for n in names}
with zipfile.ZipFile(p/'GeNeDis2026_RA_results.pptx') as z:
 embedded={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('ppt/media/')}
 assert embedded==expected,'Presentation contains missing or stale figure images'
checks.append('PPTX contains exactly the six final PNGs (SHA-256 checked)')
prs=Presentation(p/'GeNeDis2026_RA_results.pptx');assert len(prs.slides)==12
for j,s in enumerate(prs.slides):
 for sh in s.shapes:
  assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+2000 and sh.top+sh.height<=prs.slide_height+2000,(j,sh.name,'out of page')
checks.append('All 12 slides have objects within the slide bounds')
slidepdf=PdfReader(p/'GeNeDis2026_RA_results.pdf');assert len(slidepdf.pages)==12
report=PdfReader(p/'GeNeDis2026_submission_report.pdf');text='\n'.join(x.extract_text() for x in report.pages)
assert 'xwu76@uiowa.edu' in text and 'xwu76@uiowa.eduX' not in text,'Contact text is corrupt'
assert 'Full isolated analysis is in progress' not in text and 'FINAL_CLASSIFIER_RESULTS' not in text,'Unfinished report marker'
assert len(PdfReader(p/'GeNeDis2026_figure_atlas.pdf').pages)==6
for page in report.pages:
 assert abs(float(page.mediabox.width)-595.28)<1 and abs(float(page.mediabox.height)-841.89)<1,'Report page is not A4'
assert len(report.pages)==5,'Main technical report must be five pages'
assert 'Revision traceability' not in text and 'early step-36 branch' not in text,'Internal audit leaked into main report'
assert '0.900' in text and '0.979' in text and '0.2475' in text and '0.5149' in text,'Key classifier result missing'
methods=PdfReader(p/'GeNeDis2026_methods_QC.pdf');assert len(methods.pages)==5
mt='\n'.join(x.extract_text() for x in methods.pages)
assert 'xwu76@uiowa.edu' in mt and '199' in mt and '464' in mt
for page in methods.pages:
 assert abs(float(page.mediabox.width)-595.28)<1 and abs(float(page.mediabox.height)-841.89)<1,'Methods appendix page is not A4'
checks.append('Five-page technical report and five-page Methods/QC appendix; correct contact; key model results; no internal audit in main narrative; separate six-page figure atlas')
for n in names:
 im=Image.open(fig/(n+'.png'));assert min(im.info.get('dpi',(0,0)))>590,(n,'not 600 dpi')
 assert len(PdfReader(fig/(n+'.pdf')).pages)==1
checks.append('Six final figure PDFs and 600-dpi PNGs are readable')
res={'status':'PASS','checks':checks,'figure_png_sha256':{n:hashlib.sha256((fig/(n+'.png')).read_bytes()).hexdigest() for n in names}}
(r/'results/artifact_checks.json').write_text(json.dumps(res,indent=2));print(json.dumps(res,indent=2))
