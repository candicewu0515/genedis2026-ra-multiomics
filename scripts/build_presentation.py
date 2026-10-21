"""Build the submission slides from the final figures and result tables."""
from pathlib import Path
import sys,json
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
root=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else Path(__file__).resolve().parents[1]
out=root/'presentation';out.mkdir(exist_ok=True)
fig=root/'results/figures/main'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
HEADING='202020';EMPHASIS='404040';INK='202020';GREY='606060'
def box(sl,x,y,w,h,text,size=18,color=INK,bold=False):
 tb=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=tb.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(14)
 return tb
def slide(title,section):
 sl=prs.slides.add_slide(prs.slide_layouts[6])
 box(sl,.35,.16,12.6,.5,title,25,HEADING,True)
 box(sl,.35,7.13,11,.22,'GeNeDis 2026  |  Xia Wu · University of Iowa  |  '+section,10,GREY)
 box(sl,12.3,7.1,.6,.3,str(len(prs.slides)),10,GREY)
 return sl
def figure(name,title,notes):
 sl=slide(title,name)
 box(sl,.45,1.05,4.6,5.8,'\n'.join(notes),18)
 p=fig/(name+'.png');im=Image.open(p);ratio=im.width/im.height
 h=6.12;w=h*ratio
 if w>7.7:w=7.7;h=w/ratio
 sl.shapes.add_picture(str(p),Inches(5.25+(7.7-w)/2), Inches(.9+(6.12-h)/2),width=Inches(w),height=Inches(h))
 return sl
sl=slide('RA multi-omics evidence map','Individual submission')
box(sl,.6,1.25,11.8,1.55,'Genetic association, macrophage states\nand the limits of evidence in public datasets',32,INK,True)
box(sl,.65,3.35,11.6,1.1,'Xia Wu\nPhD student, University of Iowa\nxwu76@uiowa.edu',22,HEADING)
box(sl,.65,5.65,11.5,.9,'Open data · Reproducible workflow · Nature-style data figures\nExploratory RA integration with an NSCLC extension',18,GREY)
sl=slide('The scores used in this study','Score definitions')
entries=[
 ('S1 · 15 discovery genes','From sorted active/remission macrophages. Used for exploratory transfer and anti-TNF change.'),
 ('S2 · 10 predefined IFN genes','Associations in spatial tissue and across bulk disease stages. Only 2 genes overlap with S1.'),
 ('Hallmark IFN-gamma','Pathway and protein evidence. Specificity is supported in lining, not in sublining.'),
 ('M-6 / Tph · 30 genes each','Programmes from atlas states, transferred to the fluid and lung cohorts.'),
]
for j,(label,text) in enumerate(entries):
 y=1.05+j*1.35;box(sl,.6,y,4.0,.7,label,22,HEADING,True);box(sl,4.7,y,7.9,.9,text,20)
figure('Fig1','Published atlas states and tested integration baselines',[
 '31,520 nuclei from 12 donors. RNA and ATAC are measured in the same nuclei.',
 'Fine-state labels, including the 65-nucleus M-6, are taken from Gupta et al.',
 'Concatenation and WNN rank highest in this benchmark. Rare M-6 neighbourhoods favour RNA alone.'])
figure('Fig2','Genetic enrichment and spatial neighbourhood associations',[
 'Tph/Treg-associated chromatin is enriched for RA heritability.',
 'S-LDSC and scDRS are complementary methods that use the same GWAS traits.',
 'More B cells sit near CXCL13-positive T cells in 4/4 sections.',
 'The empirical global-permutation P does not account for label clustering or tissue compartments.'])
figure('Fig3','Separate discovery, tissue and treatment evidence',[
 'S1 maps to M-6. Its donor inflammation and bulk RA/OA validation tests are negative.',
 'S2 shows spatial and bulk disease-stage associations.',
 'Protein points and test use the same donor averages: 6 RA versus 3 healthy.',
 'S1 falls in good responders, but the change is not significantly different from nonresponders.'])
figure('Fig4','Candidate regulators and expression-associated genetic estimates',[
 'TF screen: two donors with M-6 and BH-adjusted gene-label permutations. Two-donor uncertainty remains.',
 'IRF5 colocalisation supports a shared association signal.',
 'Wald ratios cannot estimate drug efficacy or rule out pleiotropy.',
 'Abatacept binds CD80/CD86. IL6ST is not IL6R, the target of tocilizumab.'])
figure('ExtData_Fig1','Baselines and boundary findings',[
 'Classifiers fit all preprocessing on training data only and use repeated nested CV.',
 '199 label permutations repeat the same tuning and resampling as the observed model.',
 'Random controls match the number of retained prioritised genes in each training fold.',
 'Limits in blood transfer, response prediction, deconvolution and trajectory are all reported.'])
figure('ExtData_Fig2','Exploratory transfer to synovial fluid and NSCLC',[
 'Fluid: 9 paired patients, two therapies pooled. No response labels are available.',
 'The M-6-like fraction falls in 7/9 patients, but the paired pseudobulk-score test is not significant.',
 'LuCA: the M-6 score is higher in tumour in 20/29 paired donors.',
 'Tph coupling P = 0.092. The primary TCGA survival tests are negative.'])
sl=slide('Internal classifiers: final held-out performance','Nested CV')
df=pd.read_csv(root/'results/tables/classifier_baselines.tsv',sep='\t')
names=['F0 covariates','F1 eight programme scores','F2 top-30 prioritised genes','F3 genome-wide elastic net']
subtitle=box(sl,.6,.9,12,.75,'Mean AUC across 25 outer folds. All fitted transformations stay inside the training folds.\nF0 covariates · F1 programmes · F2 prioritised genes · F3 genome-wide',16,GREY)
for paragraph in subtitle.text_frame.paragraphs:paragraph.space_after=Pt(2)
for i,(task,g) in enumerate(df.groupby('task',sort=False)):
 x=.65+i*6.3;box(sl,x,1.7,5.7,.5,task[3:],22,HEADING,True)
 for j,n in enumerate(names):
  row=g[g.model==n].iloc[0];box(sl,x,2.45+j*.62,5.7,.5,f'{n[:2]}  {row.auc_mean:.3f} ± {row.auc_sd:.3f}',21)
 note=g[g.model.str.startswith('F1 label')]['note'].iloc[0].split(';')[0]
 random_note=g[g.model.str.startswith('F4')]['note'].iloc[0].split(';')[0]
 box(sl,x,5.1,5.7,1.1,note+'\n'+random_note,16,EMPHASIS)
box(sl,.65,6.3,12,.5,'SD across correlated folds is not a confidence interval. No external clinical validation.',18,GREY)
sl=slide('What the data support','Interpretation')
box(sl,.65,1.15,12,5.5,'A reproducible map of complementary RA associations and candidate targets for follow-up.\nScore definitions, cohort independence and strength of evidence are kept separate.\nPublished states and original discoveries are clearly credited to their sources.\nCausal regulation, therapeutic efficacy and clinical prediction need further validation.',26)
sl=slide('Reproducibility and submission contents','Validation')
box(sl,.65,1.05,12,5.9,'The full pipeline was rerun from raw inputs in a separate project. Run logs and a comparison report come with the submission.\nSaved copies of the API/GMT/eQTL inputs keep the reference version, including Open Targets 26.09.\nA quick check verifies hashes, reruns three small analyses and draws six figures without the large raw/interim data.\nREADME · Sources · Code · Processed tables · QC report · Limits · Presentation · Manifest',23)
sl=slide('Sources and contact','Provenance')
box(sl,.65,1.1,12,5.7,'Gupta et al., Nature Genetics 2023: GSE243917 atlas and author-defined states.\nAlivernini et al., Nature Medicine 2020: E-MTAB-8322 discovery cohort.\nGSE198520 / GSE89408: independent bulk cohorts.\nGSE347958 / GSE347797 / PXD062696: shared spatial/protein cohort.\nGSE296117 / LuCA / TCGA: exploratory extensions.\nAll accessions and licences: docs/sources.tsv\nXia Wu · PhD student, University of Iowa · xwu76@uiowa.edu',20)
import datetime
cp=prs.core_properties
cp.author='Xia Wu';cp.last_modified_by='Xia Wu';cp.title='RA multi-omics evidence map (GeNeDis 2026)'
cp.subject='GeNeDis 2026 Hackathon submission';cp.comments='';cp.keywords='';cp.category='';cp.revision=1
cp.created=cp.modified=datetime.datetime(2026,10,20,12,0,0)
deck=out/'GeNeDis2026_RA_results.pptx'
prs.save(deck)
# Set the editable theme colours to neutral black, white and grey. Embedded figure images keep their original bytes.
from lxml import etree
from zipfile import ZipFile, ZIP_DEFLATED
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
palette={'dk1':'000000','lt1':'FFFFFF','dk2':'202020','lt2':'F0F0F0',
         'accent1':'303030','accent2':'606060','accent3':'808080',
         'accent4':'A0A0A0','accent5':'404040','accent6':'C0C0C0',
         'hlink':'404040','folHlink':'606060'}
temporary=deck.with_suffix('.tmp.pptx')
with ZipFile(deck) as source, ZipFile(temporary,'w',ZIP_DEFLATED) as target:
 for info in source.infolist():
  data=source.read(info.filename)
  if info.filename.startswith('ppt/theme/') and info.filename.endswith('.xml'):
   xml=etree.fromstring(data)
   scheme=xml.find('.//a:clrScheme',ns)
   if scheme is not None:
    scheme.set('name','Neutral Academic')
    for slot,rgb in palette.items():
     node=scheme.find('a:'+slot,ns)
     if node is not None:
      for child in list(node):node.remove(child)
      etree.SubElement(node,'{'+ns['a']+'}srgbClr',val=rgb)
   data=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
  target.writestr(info,data)
temporary.replace(deck)
print('Saved',len(prs.slides),'slides')
